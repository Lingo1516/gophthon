


import io
import streamlit as st
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from lxml import etree
from pptx.oxml.ns import qn

# ─────────────────────────────────────────
#  色盤
# ─────────────────────────────────────────
PALETTES = {
    "🌊 Midnight（深藍）":   {"primary": "1E2761", "secondary": "CADCFC", "accent": "4FC3F7", "light": "F0F4FF",  "dark_text": "1E2761", "body_text": "374151"},
    "🧱 Terracotta（磚紅）": {"primary": "B85042", "secondary": "E7E8D1", "accent": "A7BEAE", "light": "FAF8F4",  "dark_text": "2D2D2D", "body_text": "4A5568"},
    "🌿 Forest（森林綠）":   {"primary": "2C5F2D", "secondary": "97BC62", "accent": "F5F5F5", "light": "F4FAF4",  "dark_text": "1A3A1B", "body_text": "374151"},
    "🍒 Cherry（深紅）":     {"primary": "990011", "secondary": "FCF6F5", "accent": "2F3C7E", "light": "FFF8F8",  "dark_text": "1A0005", "body_text": "4A2030"},
    "🩶 Charcoal（炭灰）":   {"primary": "36454F", "secondary": "F2F2F2", "accent": "607D8B", "light": "FAFAFA",  "dark_text": "212121", "body_text": "546E7A"},
    "🩵 Teal（青藍）":       {"primary": "028090", "secondary": "E0F7FA", "accent": "02C39A", "light": "F0FDFC",  "dark_text": "01404A", "body_text": "374151"},
}

SLIDE_TYPES = {
    "封面頁":       "cover",
    "統計數字":     "stats",
    "三欄卡片":     "three_cards",
    "左色塊清單":   "left_panel",
    "左右雙欄對比": "two_column",
    "2×2 卡片格":   "grid_2x2",
    "長條圖":       "bar_chart",
    "編號清單":     "numbered_list",
    "結尾頁":       "closing",
}

# ─────────────────────────────────────────
#  核心工具
# ─────────────────────────────────────────
def rgb(h):
    h = h.lstrip("#")
    return RGBColor(int(h[0:2],16), int(h[2:4],16), int(h[4:6],16))

def set_bg(slide, hex_color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(hex_color)

def add_rect(slide, x, y, w, h, fill_hex, line_hex=None, line_pt=0):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill_hex)
    if line_hex:
        shape.line.color.rgb = rgb(line_hex)
        shape.line.width = Pt(line_pt or 1)
    else:
        shape.line.fill.background()
    return shape

def add_oval(slide, x, y, w, h, fill_hex, alpha=0):
    shape = slide.shapes.add_shape(9, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill_hex)
    shape.line.fill.background()
    return shape

def add_text(slide, text, x, y, w, h, size=14, bold=False, color_hex="000000",
             align=PP_ALIGN.LEFT, italic=False, font="Calibri"):
    txBox = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color_hex)
    run.font.name = font
    return txBox

# ─────────────────────────────────────────
#  投影片版面
# ─────────────────────────────────────────
def slide_cover(prs, P, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, P["primary"])
    add_oval(s, 7.5,-1.3,5.0,5.0, P["accent"])
    add_oval(s,-1.2, 3.8,3.5,3.5, P["secondary"])
    add_text(s, d.get("title","標題"), 0.7,1.2,8.5,1.5, size=48, bold=True, color_hex="FFFFFF")
    add_text(s, d.get("subtitle",""), 0.7,2.85,8.0,0.65, size=19, color_hex=P["secondary"], italic=True)
    add_text(s, d.get("tag",""), 0.7,3.75,7.0,0.4, size=12, color_hex=P["secondary"])

def slide_stats(prs, P, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, P["light"])
    add_text(s, d.get("title",""), 0.5,0.25,9.0,0.65, size=30, bold=True, color_hex=P["primary"])
    add_rect(s, 0.5,1.0,9.0,0.05, P["secondary"])
    add_text(s, d.get("subtitle",""), 0.5,0.92,9.0,0.38, size=13, color_hex=P["body_text"])
    stats = d.get("stats", [])
    n = len(stats)
    card_w = 2.9 if n==3 else (4.2 if n==2 else 6.0)
    spacing = (9.0 - card_w*n) / (n+1)
    for i,(num,label) in enumerate(stats):
        x = spacing + i*(card_w+spacing) + 0.5
        add_rect(s,x,1.3,card_w,2.9,"FFFFFF", P["secondary"],1)
        add_text(s,num,x+0.1,1.65,card_w-0.2,1.1, size=42,bold=True,color_hex=P["primary"],align=PP_ALIGN.CENTER)
        add_text(s,label,x+0.1,2.9,card_w-0.2,1.0, size=13,color_hex=P["body_text"],align=PP_ALIGN.CENTER)

def slide_three_cards(prs, P, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s, P["light"])
    add_text(s, d.get("title",""), 0.5,0.25,9.0,0.65, size=30,bold=True,color_hex=P["primary"])
    add_text(s, d.get("subtitle",""), 0.5,0.92,9.0,0.38, size=13,color_hex=P["body_text"])
    for i,card in enumerate(d.get("cards",[])[:3]):
        x = 0.35+i*3.2
        add_rect(s,x,1.15,3.0,3.85,"FFFFFF",P["secondary"],1)
        add_rect(s,x,1.15,3.0,0.09,P["primary"])
        add_text(s,card.get("icon",""),x+0.1,1.38,2.8,0.6, size=26,align=PP_ALIGN.CENTER)
        add_text(s,card.get("title",""),x+0.15,2.1,2.7,0.55, size=15,bold=True,color_hex=P["primary"],align=PP_ALIGN.CENTER)
        add_text(s,card.get("body",""),x+0.15,2.72,2.7,2.0, size=12,color_hex=P["body_text"],align=PP_ALIGN.CENTER)

def slide_left_panel(prs, P, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s,"FFFFFF")
    add_rect(s,0,0,3.2,5.625,P["primary"])
    add_text(s,d.get("number","01"), 0.1,0.3,3.0,1.0, size=60,bold=True,color_hex=P["accent"],align=PP_ALIGN.CENTER)
    add_text(s,d.get("chapter_title",""), 0.1,1.5,3.0,1.8, size=28,bold=True,color_hex="FFFFFF",align=PP_ALIGN.CENTER)
    add_text(s,d.get("subtitle",""), 3.5,0.3,6.2,0.65, size=22,bold=True,color_hex=P["dark_text"])
    for i,(icon,text) in enumerate(d.get("items",[])):
        y = 1.1+i*1.0
        add_rect(s,3.5,y,6.0,0.82,P["light"],P["secondary"],1)
        add_text(s,icon, 3.6,y+0.08,0.6,0.65, size=18,align=PP_ALIGN.CENTER)
        add_text(s,text, 4.3,y+0.1,5.0,0.62, size=13,color_hex=P["dark_text"])

def slide_two_column(prs, P, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s,P["light"])
    add_text(s,d.get("title",""), 0.5,0.25,9.0,0.65, size=30,bold=True,color_hex=P["primary"])
    cols=[
        (d.get("left_title",""),  d.get("left_items",[]),  0.5, P["primary"]),
        (d.get("right_title",""), d.get("right_items",[]), 5.3, P["accent"]),
    ]
    for col_title,items,cx,bar in cols:
        add_rect(s,cx,1.4,4.2,3.8,"FFFFFF",P["secondary"],1)
        add_rect(s,cx,1.4,4.2,0.12,bar)
        add_text(s,col_title, cx+0.15,1.62,4.0,0.5, size=15,bold=True,color_hex=bar)
        for j,item in enumerate(items[:4]):
            add_text(s,f"▸  {item}", cx+0.2,2.22+j*0.65,3.9,0.55, size=13,color_hex=P["dark_text"])

def slide_grid_2x2(prs, P, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s,P["light"])
    add_text(s,d.get("title",""), 0.5,0.25,9.0,0.65, size=30,bold=True,color_hex=P["primary"])
    positions=[(0.5,1.4),(5.1,1.4),(0.5,3.2),(5.1,3.2)]
    for i,card in enumerate(d.get("cards",[])[:4]):
        x,y=positions[i]
        add_rect(s,x,y,4.1,1.7,"FFFFFF",P["secondary"],1)
        add_rect(s,x,y,0.12,1.7,P["primary"])
        add_text(s,card.get("title",""), x+0.22,y+0.12,3.8,0.48, size=15,bold=True,color_hex=P["primary"])
        add_text(s,card.get("body",""),  x+0.22,y+0.65,3.8,0.95, size=12,color_hex=P["dark_text"])

def slide_bar_chart(prs, P, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s,"FFFFFF")
    add_text(s,d.get("title",""), 0.5,0.25,9.0,0.65, size=30,bold=True,color_hex=P["primary"])
    add_text(s,d.get("subtitle",""), 0.5,0.92,9.0,0.38, size=13,color_hex=P["body_text"])
    labels = d.get("labels",[])
    values = d.get("values",[])
    if labels and values:
        cd = ChartData()
        cd.categories = labels
        cd.add_series(d.get("unit","數值"), values)
        chart = s.shapes.add_chart(
            XL_CHART_TYPE.COLUMN_CLUSTERED,
            Inches(0.5),Inches(1.35),Inches(9),Inches(3.95), cd
        ).chart
        chart.series[0].format.fill.solid()
        chart.series[0].format.fill.fore_color.rgb = rgb(P["primary"])
        chart.has_legend = False

def slide_numbered_list(prs, P, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s,P["primary"])
    add_oval(s,7.0,-1.2,4.5,4.5,P["accent"])
    add_oval(s,-1.0,3.5,3.5,3.5,P["secondary"])
    add_text(s,d.get("title",""), 0.7,0.3,8.0,0.8, size=34,bold=True,color_hex="FFFFFF")
    for i,text in enumerate(d.get("items",[])[:5]):
        y=1.25+i*0.98
        add_oval(s,0.7,y+0.08,0.55,0.55,P["secondary"])
        add_text(s,str(i+1), 0.7,y+0.08,0.55,0.55, size=16,bold=True,color_hex=P["primary"],align=PP_ALIGN.CENTER)
        add_text(s,text, 1.45,y+0.1,7.8,0.55, size=14,color_hex="FFFFFF")

def slide_closing(prs, P, d):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(s,P["primary"])
    add_oval(s,7.5,-1.3,5.0,5.0,P["accent"])
    add_oval(s,-1.2,3.8,3.5,3.5,P["secondary"])
    add_text(s,d.get("headline",""), 1.0,1.0,8.0,2.8, size=40,bold=True,color_hex="FFFFFF",align=PP_ALIGN.CENTER)
    add_rect(s,3.8,3.9,2.4,0.06,P["secondary"])
    add_text(s,d.get("tagline",""), 1.0,4.1,8.0,0.5, size=14,color_hex=P["secondary"],align=PP_ALIGN.CENTER,italic=True)

SLIDE_BUILDERS = {
    "cover":        slide_cover,
    "stats":        slide_stats,
    "three_cards":  slide_three_cards,
    "left_panel":   slide_left_panel,
    "two_column":   slide_two_column,
    "grid_2x2":     slide_grid_2x2,
    "bar_chart":    slide_bar_chart,
    "numbered_list":slide_numbered_list,
    "closing":      slide_closing,
}

def build_pptx(theme_name, slides_data, filename="簡報.pptx"):
    P = PALETTES[theme_name]
    prs = Presentation()
    prs.slide_width  = Inches(10)
    prs.slide_height = Inches(5.625)
    for slide_type, data in slides_data:
        builder = SLIDE_BUILDERS.get(slide_type)
        if builder:
            builder(prs, P, data)
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf

# ─────────────────────────────────────────
#  Streamlit UI
# ─────────────────────────────────────────
st.set_page_config(page_title="簡報生成器", page_icon="📊", layout="wide")

st.title("📊 簡報生成器")
st.caption("填入內容，一鍵產生專業 .pptx 簡報")

# ── 側邊欄：主題 & 檔名 ──
with st.sidebar:
    st.header("⚙️ 設定")
    theme = st.selectbox("主題色系", list(PALETTES.keys()))
    filename = st.text_input("檔案名稱", "我的簡報.pptx")
    st.divider()
    st.caption("版面說明")
    for k in SLIDE_TYPES:
        st.markdown(f"• **{k}**")

# ── 主頁：投影片編輯 ──
st.subheader("投影片設定")

if "slides" not in st.session_state:
    st.session_state.slides = [
        {"type": "封面頁",   "data": {"title":"簡報標題", "subtitle":"副標題說明", "tag":"部門 · 日期"}},
        {"type": "統計數字", "data": {"title":"關鍵數字", "subtitle":"", "stats":[("72%","說明一"),("$1.8T","說明二"),("40%","說明三")]}},
        {"type": "結尾頁",   "data": {"headline":"感謝您的聆聽", "tagline":"如有問題歡迎討論"}},
    ]

slides_data = []
to_delete = None

for idx, slide in enumerate(st.session_state.slides):
    with st.expander(f"第 {idx+1} 頁 — {slide['type']}", expanded=(idx==0)):
        col1, col2 = st.columns([3,1])
        with col2:
            if st.button("🗑️ 刪除", key=f"del_{idx}"):
                to_delete = idx
        with col1:
            stype = st.selectbox("版面類型", list(SLIDE_TYPES.keys()),
                                 index=list(SLIDE_TYPES.keys()).index(slide["type"]),
                                 key=f"type_{idx}")
            st.session_state.slides[idx]["type"] = stype

        d = slide["data"]
        t = SLIDE_TYPES[stype]

        # 依版面類型顯示不同欄位
        if t == "cover":
            d["title"]    = st.text_input("大標題", d.get("title",""), key=f"title_{idx}")
            d["subtitle"] = st.text_input("副標題", d.get("subtitle",""), key=f"sub_{idx}")
            d["tag"]      = st.text_input("標籤列（部門/日期）", d.get("tag",""), key=f"tag_{idx}")

        elif t == "stats":
            d["title"]    = st.text_input("標題", d.get("title",""), key=f"title_{idx}")
            d["subtitle"] = st.text_input("副標題", d.get("subtitle",""), key=f"sub_{idx}")
            st.markdown("**統計數字（最多3個）**")
            stats = d.get("stats", [("",""),("",""),("","")])
            new_stats = []
            for si in range(3):
                c1,c2 = st.columns(2)
                num   = c1.text_input(f"數字 {si+1}", stats[si][0] if si<len(stats) else "", key=f"num_{idx}_{si}")
                label = c2.text_input(f"說明 {si+1}", stats[si][1] if si<len(stats) else "", key=f"lbl_{idx}_{si}")
                if num:
                    new_stats.append((num, label))
            d["stats"] = new_stats

        elif t == "three_cards":
            d["title"]    = st.text_input("標題", d.get("title",""), key=f"title_{idx}")
            d["subtitle"] = st.text_input("副標題", d.get("subtitle",""), key=f"sub_{idx}")
            cards = d.get("cards",[{},{},{}])
            new_cards = []
            for ci in range(3):
                st.markdown(f"**卡片 {ci+1}**")
                c1,c2,c3 = st.columns(3)
                icon  = c1.text_input("Icon", cards[ci].get("icon","") if ci<len(cards) else "", key=f"icon_{idx}_{ci}")
                title = c2.text_input("標題", cards[ci].get("title","") if ci<len(cards) else "", key=f"ctitle_{idx}_{ci}")
                body  = c3.text_input("說明", cards[ci].get("body","") if ci<len(cards) else "", key=f"cbody_{idx}_{ci}")
                new_cards.append({"icon":icon,"title":title,"body":body})
            d["cards"] = new_cards

        elif t == "left_panel":
            d["number"]        = st.text_input("章節號碼", d.get("number","01"), key=f"num_{idx}")
            d["chapter_title"] = st.text_input("左側標題（換行用\\n）", d.get("chapter_title",""), key=f"ch_{idx}").replace("\\n","\n")
            d["subtitle"]      = st.text_input("右側大標", d.get("subtitle",""), key=f"sub_{idx}")
            st.markdown("**清單項目（最多4個）**")
            items = d.get("items",[])
            new_items = []
            for ii in range(4):
                c1,c2 = st.columns([1,5])
                icon = c1.text_input("Icon", items[ii][0] if ii<len(items) else "", key=f"icon_{idx}_{ii}")
                text = c2.text_input("文字", items[ii][1] if ii<len(items) else "", key=f"itext_{idx}_{ii}")
                if text:
                    new_items.append((icon, text))
            d["items"] = new_items

        elif t == "two_column":
            d["title"]       = st.text_input("標題", d.get("title",""), key=f"title_{idx}")
            c1,c2 = st.columns(2)
            d["left_title"]  = c1.text_input("左欄標題", d.get("left_title",""), key=f"lt_{idx}")
            d["right_title"] = c2.text_input("右欄標題", d.get("right_title",""), key=f"rt_{idx}")
            left_raw  = c1.text_area("左欄項目（每行一項）", "\n".join(d.get("left_items",[])),  key=f"li_{idx}")
            right_raw = c2.text_area("右欄項目（每行一項）", "\n".join(d.get("right_items",[])), key=f"ri_{idx}")
            d["left_items"]  = [x for x in left_raw.split("\n") if x.strip()]
            d["right_items"] = [x for x in right_raw.split("\n") if x.strip()]

        elif t == "grid_2x2":
            d["title"] = st.text_input("標題", d.get("title",""), key=f"title_{idx}")
            cards = d.get("cards",[{},{},{},{}])
            new_cards = []
            for ci in range(4):
                c1,c2 = st.columns(2)
                title = c1.text_input(f"卡片{ci+1} 標題", cards[ci].get("title","") if ci<len(cards) else "", key=f"gt_{idx}_{ci}")
                body  = c2.text_input(f"卡片{ci+1} 說明", cards[ci].get("body","") if ci<len(cards) else "",  key=f"gb_{idx}_{ci}")
                new_cards.append({"title":title,"body":body})
            d["cards"] = new_cards

        elif t == "bar_chart":
            d["title"]    = st.text_input("標題", d.get("title",""), key=f"title_{idx}")
            d["subtitle"] = st.text_input("副標題", d.get("subtitle",""), key=f"sub_{idx}")
            d["unit"]     = st.text_input("數列名稱", d.get("unit","數值"), key=f"unit_{idx}")
            labels_raw = st.text_input("X 軸標籤（逗號分隔）", ",".join(d.get("labels",[])), key=f"labels_{idx}")
            values_raw = st.text_input("數值（逗號分隔）",      ",".join(map(str,d.get("values",[]))), key=f"values_{idx}")
            d["labels"] = [x.strip() for x in labels_raw.split(",") if x.strip()]
            try:
                d["values"] = [float(x.strip()) for x in values_raw.split(",") if x.strip()]
            except:
                d["values"] = []

        elif t == "numbered_list":
            d["title"] = st.text_input("標題", d.get("title",""), key=f"title_{idx}")
            items_raw  = st.text_area("項目（每行一項，最多5項）", "\n".join(d.get("items",[])), key=f"items_{idx}")
            d["items"] = [x for x in items_raw.split("\n") if x.strip()][:5]

        elif t == "closing":
            d["headline"] = st.text_input("結語標題（換行用\\n）", d.get("headline","").replace("\n","\\n"), key=f"hl_{idx}").replace("\\n","\n")
            d["tagline"]  = st.text_input("標語", d.get("tagline",""), key=f"tl_{idx}")

        st.session_state.slides[idx]["data"] = d
        slides_data.append((SLIDE_TYPES[stype], d))

if to_delete is not None:
    st.session_state.slides.pop(to_delete)
    st.rerun()

# ── 新增投影片 ──
st.divider()
col1, col2 = st.columns([3,1])
with col1:
    new_type = st.selectbox("新增版面", list(SLIDE_TYPES.keys()), key="new_type")
with col2:
    if st.button("➕ 新增投影片", use_container_width=True):
        st.session_state.slides.append({"type": new_type, "data": {}})
        st.rerun()

# ── 生成按鈕 ──
st.divider()
if st.button("🚀 生成簡報", type="primary", use_container_width=True):
    with st.spinner("產生中..."):
        try:
            buf = build_pptx(theme, slides_data, filename)
            st.success(f"✅ 完成！共 {len(slides_data)} 頁")
            st.download_button(
                label="⬇️ 下載 .pptx",
                data=buf,
                file_name=filename,
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                use_container_width=True,
            )
        except Exception as e:
            st.error(f"❌ 錯誤：{e}")
